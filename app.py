import streamlit as st
import pandas as pd
import xlrd
import openpyxl
import random
import io
import zipfile
import gc
import time
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from collections import Counter

# =========================================================
#  🎨 إعدادات الصفحة
# =========================================================
st.set_page_config(page_title="نظام الاختبارات الذكي", page_icon="🎓", layout="wide")

st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;700&display=swap');
    * {font-family: 'Cairo', sans-serif;}
    .main {direction: rtl;}
    .stButton button {width: 100%; font-weight: bold; font-size: 18px; padding: 10px; background-color: #ff4b4b; color: white;}
    div[data-testid="stExpander"] {border: 1px solid #ddd; border-radius: 8px;}
    .log-box {background-color: #262730; color: #00ff00; padding: 10px; border-radius: 5px; height: 200px; overflow-y: auto; direction: ltr; text-align: left; font-family: monospace;}
</style>
""", unsafe_allow_html=True)

# =========================================================
#  ⚙️ المنطق (القراءة الذكية + التوليد)
# =========================================================
def normalize_text(text):
    if pd.isna(text) or str(text).strip() == "": return ""
    text = str(text).strip()
    return text.replace('ة', 'ه').replace('ى', 'ي').replace('أ', 'ا').replace('إ', 'ا').replace('آ', 'ا')

def clean_for_comp(text): return normalize_text(text).replace(" ", "")

def force_align_options(options, correct_text, target_idx):
    final_opts = list(options) if options else []
    while len(final_opts) < 4: final_opts.append("---")
    current_idx = -1
    clean_corr = clean_for_comp(correct_text)
    for i, opt in enumerate(final_opts):
        if clean_for_comp(str(opt)) == clean_corr: current_idx = i; break
    if current_idx != -1:
        if current_idx != target_idx:
            final_opts[target_idx], final_opts[current_idx] = final_opts[current_idx], final_opts[target_idx]
    else: final_opts[target_idx] = correct_text
    return final_opts[:4]

# --- Word Helpers ---
def set_section_rtl(section):
    if not section._sectPr.find(qn('w:bidi')):
        section._sectPr.append(OxmlElement('w:bidi'))
    section.left_margin = Inches(0.4); section.right_margin = Inches(0.4)

def fix_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.bidi = True
    p.paragraph_format.space_after = Pt(4)

def format_table(table):
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None: tbl_pr = table._tbl.add_tblPr()
    if not tbl_pr.find(qn('w:bidiVisual')): tbl_pr.append(OxmlElement('w:bidiVisual'))
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'right'); tbl_pr.append(jc)
    tbl_borders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement(f'w:{b}'); el.set(qn('w:val'), 'nil'); tbl_borders.append(el)
    tbl_pr.append(tbl_borders)

def set_font(run, size=14, bold=False):
    run.font.name = 'Times New Roman'; run.font.size = Pt(size); run.bold = bold
    rPr = run._element.get_or_add_rPr()
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Times New Roman'); fonts.set(qn('w:hAnsi'), 'Times New Roman')
    fonts.set(qn('w:eastAsia'), 'Times New Roman'); fonts.set(qn('w:cs'), 'Times New Roman')
    rPr.append(fonts)
    if not rPr.find(qn('w:rtl')): rPr.append(OxmlElement('w:rtl'))

def add_question(doc, num, text, opts):
    try:
        t = doc.add_table(rows=2, cols=4); format_table(t)
        c = t.cell(0, 0).merge(t.cell(0, 3))
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.bidi = True
        set_font(p.add_run(f"{num}) {text}"), 14, True)
        chars = ['أ', 'ب', 'ج', 'د']
        for i, opt in enumerate(opts):
            if i < 4:
                c2 = t.cell(1, i); p2 = c2.paragraphs[0]
                p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p2.paragraph_format.bidi = True
                set_font(p2.add_run(f"{chars[i] if i<len(chars) else '-'}. {opt}"), 14, False)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
    except: pass

# --- القارئ الذكي جداً (Super Smart Reader) ---
# يبحث عن كلمة "سؤال" ويأخذ الأعمدة الأربعة التي تليها مباشرة كخيارات
@st.cache_data(show_spinner=False)
def read_question_bank_smart(file_bytes, filename):
    file_bytes.seek(0)
    try:
        # 1. XLSX
        if filename.lower().endswith('.xlsx'):
            wb = openpyxl.load_workbook(file_bytes, data_only=False); sh = wb.active
            rows = list(sh.iter_rows())
            hr = -1; cols = {'u': -1, 'q': -1, 'obj': -1}
            # بحث عن الهيدر
            for r_idx, row in enumerate(rows[:25]): # وسعنا النطاق لـ 25 صف
                vs = [normalize_text(c.value) for c in row]
                if any('سؤال' in x for x in vs):
                    hr = r_idx
                    for c_idx, v in enumerate(vs):
                        if 'وحده' in v: cols['u'] = c_idx
                        elif 'سؤال' in v: cols['q'] = c_idx
                        elif 'هدف' in v: cols['obj'] = c_idx
                    break
            
            if hr == -1 or cols['q'] == -1: return pd.DataFrame()
            
            # القاعدة الذهبية: الاختيارات هي الـ 4 أعمدة بعد السؤال مباشرة
            s_opt = cols['q'] + 1
            opt_cols = [s_opt, s_opt+1, s_opt+2, s_opt+3]
            
            data = []
            for row in rows[hr+1:]:
                try: q = str(row[cols['q']].value if row[cols['q']].value else "").strip()
                except: continue
                if not q: continue
                
                u = "عام"
                if cols['u'] != -1 and cols['u'] < len(row): 
                    val = row[cols['u']].value; u = normalize_text(val) if val else "عام"
                cat = u + (str(row[cols['obj']].value) if cols['obj']!=-1 and cols['obj']<len(row) and row[cols['obj']].value else "")
                
                o_txt = []; o_is_col = []
                for ci in opt_cols:
                    if ci < len(row):
                        c = row[ci]; val = str(c.value if c.value else "").strip()
                        is_col = False
                        if c.fill and c.fill.start_color:
                            if c.fill.start_color.type == 'rgb' and c.fill.start_color.rgb not in ['00000000', 'FFFFFFFF', None]: is_col=True
                            elif c.fill.start_color.type == 'theme': is_col=True
                        o_txt.append(val); o_is_col.append(is_col)
                    else: o_txt.append(""); o_is_col.append(False)
                
                corr = ""; found_idx = -1
                for i in range(len(o_txt)):
                    if o_txt[i] and o_is_col[i]: found_idx = i; break
                
                if found_idx != -1: corr = o_txt[found_idx]
                real = [x for x in o_txt if x]
                if real and corr:
                    data.append({'category': cat, 'question':q, 'options':real[:4], 'correct_text':corr})
            return pd.DataFrame(data)

        # 2. XLS
        elif filename.lower().endswith('.xls'):
            book = xlrd.open_workbook(file_contents=file_bytes.read(), formatting_info=True)
            sh = book.sheet_by_index(0)
            for n in book.sheet_names():
                if 'بنك' in n or 'اسئله' in normalize_text(n): sh = book.sheet_by_name(n); break
            
            hr = -1; cols = {'u': -1, 'q': -1, 'obj': -1}
            for r in range(min(25, sh.nrows)):
                vs = [normalize_text(sh.cell_value(r, c)) for c in range(sh.ncols)]
                if any('سؤال' in x for x in vs):
                    hr = r
                    for c, v in enumerate(vs):
                        if 'وحده' in v: cols['u'] = c
                        elif 'سؤال' in v: cols['q'] = c
                        elif 'هدف' in v: cols['obj'] = c
                    break
            
            if hr == -1 or cols['q'] == -1: return pd.DataFrame()
            
            s_opt = cols['q'] + 1
            opt_cols = [s_opt, s_opt+1, s_opt+2, s_opt+3]
            
            data = []
            for r in range(hr+1, sh.nrows):
                q = str(sh.cell_value(r, cols['q'])).strip()
                if not q: continue
                u = str(sh.cell_value(r, cols['u'])).strip() if cols['u']!=-1 else "عام"
                cat = u + (str(sh.cell_value(r, cols['obj'])) if cols['obj']!=-1 else "")

                o_txt = []; o_clr = []
                for ci in opt_cols:
                    if ci < sh.ncols:
                        o_txt.append(str(sh.cell_value(r, ci)).strip())
                        o_clr.append(book.xf_list[sh.cell_xf_index(r, ci)].background.pattern_colour_index)
                    else: o_txt.append(""); o_clr.append(64)
                
                corr = ""; found_idx = -1
                valid_clrs = [o_clr[i] for i, txt in enumerate(o_txt) if txt]
                if valid_clrs:
                    cnt = Counter(valid_clrs)
                    uniq = next((k for k,v in cnt.items() if v==1), None)
                    if uniq: 
                        for i in range(len(o_txt)):
                            if o_txt[i] and o_clr[i] == uniq: found_idx = i; break
                    else:
                        for i in range(len(o_txt)):
                            if o_txt[i] and o_clr[i] != 64: found_idx = i; break
                
                if found_idx != -1: corr = o_txt[found_idx]
                real = [x for x in o_txt if x]
                if real and corr:
                    data.append({'category': cat, 'question':q, 'options':real[:4], 'correct_text':corr})
            return pd.DataFrame(data)

    except: return pd.DataFrame()
    return pd.DataFrame()

def get_master_pattern(file_obj, limit=30):
    if not file_obj: return [random.randint(0,3) for _ in range(limit)]
    file_obj.seek(0)
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=False); sh = wb.active
        pat = []
        for row in sh.iter_rows():
            found = -1
            for cell in row:
                is_col = False
                if cell.fill and cell.fill.start_color:
                    if cell.fill.start_color.type == 'rgb' and cell.fill.start_color.rgb not in ['00000000', 'FFFFFFFF', None]: is_col=True
                    elif cell.fill.start_color.type == 'theme': is_col=True
                val = str(cell.value).strip() if cell.value else ""
                if (is_col and val) or ('*' in val):
                    if 'أ' in val or 'ا' in val: found=0
                    elif 'ب' in val: found=1
                    elif 'ج' in val: found=2
                    elif 'د' in val: found=3
                    if found!=-1: pat.append(found); break
        while len(pat) < limit: pat.append(random.randint(0,3))
        return pat
    except: return [random.randint(0,3) for _ in range(limit)]

def generate_exam(df, total):
    if df.empty: return pd.DataFrame()
    grp = df.groupby('category'); cats = list(grp.groups.keys())
    base = total // len(cats) if cats else 0; rem = total % len(cats) if cats else 0
    sel = []; random.shuffle(cats)
    for i, cat in enumerate(cats):
        q = base + (1 if i < rem else 0)
        g = grp.get_group(cat)
        sel.append(g.sample(n=q) if len(g)>=q else g)
    res = pd.concat(sel) if sel else pd.DataFrame()
    if len(res) < total:
        left = df[~df.index.isin(res.index)]
        if not left.empty: res = pd.concat([res, left.sample(n=total-len(res)) if len(left)>=total-len(res) else left])
    return res.sample(frac=1).reset_index(drop=True)

# =========================================================
#  🖥️ واجهة المستخدم
# =========================================================

st.title("🚀 نظام توليد الاختبارات (Enterprise Edition)")
st.info("قم برفع بنوك الأسئلة، ثم حدد الإعدادات لكل نموذج. يدعم النظام آلاف الأسئلة.")

col1, col2 = st.columns([1, 2])
with col1:
    st.markdown("#### ⚙️ التحكم")
    num_questions = st.number_input("عدد الأسئلة في النموذج", min_value=5, value=30)
    st.markdown("#### 📜 السجل الحي")
    log_box = st.empty()
    logs = []
    def log(msg):
        logs.append(f"{time.strftime('%H:%M:%S')} | {msg}")
        log_box.markdown(f'<div class="log-box">{"<br>".join(logs[-12:])}</div>', unsafe_allow_html=True)

with col2:
    st.markdown("#### 1️⃣ بنوك الأسئلة (Excel)")
    uploaded_banks = st.file_uploader("ارفع ملفات (xls/xlsx)", type=['xls', 'xlsx'], accept_multiple_files=True)

st.markdown("---")
st.markdown("#### 2️⃣ إعدادات النماذج")

models_conf = [
    {"id": "1", "name": "ا_صباحي", "folder": "صباحي"},
    {"id": "2", "name": "ب_صباحي", "folder": "صباحي"},
    {"id": "3", "name": "ا_مسائي", "folder": "مسائي"},
    {"id": "4", "name": "ب_مسائي", "folder": "مسائي"},
    {"id": "5", "name": "ا_دور_ثاني", "folder": "دور_ثاني"},
    {"id": "6", "name": "ب_دور_ثاني", "folder": "دور_ثاني"},
]
model_files = {}
cols = st.columns(3)
for i, m in enumerate(models_conf):
    with cols[i % 3]:
        with st.expander(f"📌 {m['name']}", expanded=True):
            t = st.file_uploader("التمبلت (Word)", type=['docx'], key=f"t_{m['id']}")
            k = st.file_uploader("المفتاح (Excel ملون)", type=['xlsx'], key=f"k_{m['id']}")
            model_files[m['name']] = {"folder": m['folder'], "t": t, "k": k}

st.markdown("---")
if st.button("🚀 إنشاء وتحميل الملفات", use_container_width=True):
    if not uploaded_banks:
        st.error("⚠️ يرجى رفع بنوك الأسئلة!")
    else:
        try:
            log("بدء المعالجة...")
            all_dfs = []
            progress = st.progress(0)
            
            for i, f in enumerate(uploaded_banks):
                log(f"قراءة: {f.name}")
                df = read_question_bank_smart(f, f.name)
                if not df.empty: 
                    all_dfs.append(df)
                    log(f"✅ تم قراءة {len(df)} سؤال.")
                else:
                    log(f"⚠️ تحذير: لم يتم استخراج أسئلة من {f.name}")
                progress.progress((i + 1) / len(uploaded_banks))
            
            if not all_dfs:
                st.error("❌ لم يتم العثور على أي أسئلة صالحة! تأكد من وجود عمود 'سؤال' في الملفات.")
            else:
                BIG_DF = pd.concat(all_dfs).reset_index(drop=True)
                log(f"الإجمالي: {len(BIG_DF)} سؤال متاح.")
                
                zip_buffer = io.BytesIO()
                with zipfile.ZipFile(zip_buffer, "w") as zf:
                    for i, (m_name, conf) in enumerate(model_files.items()):
                        log(f"إنشاء: {m_name}...")
                        pattern = get_master_pattern(conf['k'], num_questions)
                        exam = generate_exam(BIG_DF, num_questions)
                        if len(exam) > num_questions: exam = exam.iloc[:num_questions]
                        exam = exam.reset_index(drop=True)
                        
                        doc = Document(conf['t']) if conf['t'] else Document()
                        set_section_rtl(doc.sections[0])
                        doc.add_paragraph("")
                        sub = doc.add_paragraph(); fix_paragraph(sub)
                        set_font(sub.add_run("اختر الإجابة الصحيحة :"), 14, True)
                        
                        for idx, row in enumerate(exam.to_dict('records')):
                            t_idx = pattern[idx % len(pattern)]
                            opts = force_align_options(row['options'], row['correct_text'], t_idx)
                            add_question(doc, idx+1, row['question'], opts)
                        
                        bio = io.BytesIO(); doc.save(bio)
                        zf.writestr(f"{conf['folder']}/{m_name}.docx", bio.getvalue())
                
                st.success("✅ تم الانتهاء بنجاح!")
                st.download_button("📥 تحميل ZIP", zip_buffer.getvalue(), "Exams_Pack.zip", "application/zip", use_container_width=True)
                gc.collect()
        except Exception as e:
            st.error(f"Error: {e}")
            log(f"Error: {e}")
